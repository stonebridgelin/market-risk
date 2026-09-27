"""导入截图时代的回测记录（docs/STORAGE.md 第7节 import-legacy）。

- 只读取录入值（分数与原始读数）；公式列由程序重新计算，并与 Excel 中的缓存值比对作为校验。
- 已完成的样本以 data_source_type=screenshot 导入为运行目录，设为正式记录（set_by=import-legacy，reviewed=true）。
- 阶段4的程序记录作为对照保留；两者差异写入 reviews；若任何维度分数不同，不自动处理，列出交给用户判断。
- "变更记录"工作表导入为复核记录（reviewer=legacy_changelog）。
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from market_risk.materials import Review, add_reviews
from market_risk.models import DimensionScore
from market_risk.scoring.common import assemble
from market_risk.storage.paths import MARKET, RISK_SCORING, RUN_FILES, StoragePaths
from market_risk.storage.runs import (
    STATUS_COMPLETE,
    GitInfo,
    base_meta,
    list_runs,
    new_run,
    read_json,
    set_official,
    write_json,
)

SUMMARY_SHEET, RAW_SHEET, CHANGELOG_SHEET = "样本汇总", "原始数据", "变更记录"
HEADER_ROW, FIRST_ROW = 4, 5
DIM_NAMES = ("价格", "广度", "VIX", "利率", "信用")
VERSIONS = {"v2-M": "EFGHI", "v3-R1": "MNOPQ"}
# 原始数据表：录入列 → 指标名（百分比以小数存储，导入时 ×100 转为百分数）
RAW_COLUMNS = {
    "C": "SPY.close", "D": "SPY.ma5", "E": "SPY.ma20", "F": "SPY.ma50", "G": "SPY.ma200",
    "H": "QQQ.close", "I": "QQQ.ma5", "J": "QQQ.ma20", "K": "QQQ.ma50", "L": "QQQ.ma200",
    "M": "RSP.close", "N": "RSP.ma5", "O": "RSP.ma20", "P": "RSP.ma50", "Q": "RSP.ma200",
    "R": "F", "S": "W", "T": "F5", "U": "W5", "V": "VIX", "W": "VIX_T5", "X": "y", "Y": "H",
    "Z": "H_date", "AA": "y_t20", "AB": "oas_o1_date", "AC": "oas_o1", "AD": "oas_o6_date", "AE": "oas_o6",
    "AF": "hyg_lqd", "AG": "data_note",
}
PERCENT_KEYS = {"F", "W", "F5", "W5", "y", "H", "y_t20", "oas_o1", "oas_o6"}
PRICE_TOL = 0.02


class LegacyError(ValueError):
    pass


@dataclass
class LegacySample:
    row: int
    seq: int
    base_date: dt.date
    status: str
    scores: dict[str, list[int | None]]
    cached: dict[str, Any]                  # 样本汇总中公式列的缓存值
    texts: dict[str, Any]                   # 结果窗口、风险事件、贴近门槛、备注、会话
    raw: dict[str, Any] = field(default_factory=dict)
    raw_cached: dict[str, Any] = field(default_factory=dict)

    @property
    def complete(self) -> bool:
        return self.status == "已完成" and all(
            s is not None for v in self.scores.values() for s in v
        )


@dataclass
class LegacyImportResult:
    imported: list[tuple[dt.date, str]] = field(default_factory=list)      # (基准日, run_id)
    skipped: list[str] = field(default_factory=list)
    formula_checks: list[str] = field(default_factory=list)                # 公式重算与缓存值不一致
    score_differences: list[str] = field(default_factory=list)             # 需用户判断
    metric_differences: list[str] = field(default_factory=list)
    changelog_rows: int = 0


def _as_date(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip()):
        return dt.date.fromisoformat(v.strip())
    return None


def read_excel(path: Path) -> tuple[list[LegacySample], list[dict[str, Any]]]:
    """读取样本汇总（录入值 + 公式缓存值）、原始数据、变更记录。"""
    if not path.exists():
        raise LegacyError(f"找不到旧 Excel：{path}")
    values = openpyxl.load_workbook(path, data_only=True)
    for sheet in (SUMMARY_SHEET, RAW_SHEET, CHANGELOG_SHEET):
        if sheet not in values.sheetnames:
            raise LegacyError(f"Excel 缺少工作表：{sheet}")
    ws = values[SUMMARY_SHEET]
    samples: list[LegacySample] = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        base = _as_date(ws[f"B{r}"].value)
        if base is None:
            continue
        scores = {ver: [ws[f"{c}{r}"].value for c in cols] for ver, cols in VERSIONS.items()}
        samples.append(LegacySample(
            row=r, seq=int(ws[f"A{r}"].value), base_date=base, status=str(ws[f"D{r}"].value or ""),
            scores={k: [None if v is None else int(v) for v in vs] for k, vs in scores.items()},
            cached={c: ws[f"{c}{r}"].value for c in ("J", "K", "L", "R", "S", "T", "U")},
            texts={"weekday": ws[f"C{r}"].value, "outcome_window": ws[f"V{r}"].value,
                   "risk_event": ws[f"W{r}"].value, "near_threshold": ws[f"X{r}"].value,
                   "remark": ws[f"Y{r}"].value, "session": ws[f"Z{r}"].value},
        ))
    raw_ws = values[RAW_SHEET]
    by_date = {s.base_date: s for s in samples}
    formula_cols = [openpyxl.utils.get_column_letter(i) for i in range(34, 51)]   # AH..AX
    for r in range(FIRST_ROW, raw_ws.max_row + 1):
        base = _as_date(raw_ws[f"B{r}"].value)
        if base is None or base not in by_date:
            continue
        s = by_date[base]
        for col, key in RAW_COLUMNS.items():
            v = raw_ws[f"{col}{r}"].value
            if isinstance(v, dt.datetime):
                v = v.date().isoformat()
            elif key in PERCENT_KEYS and v is not None:
                v = round(float(v) * 100, 6)
            s.raw[key] = v
        s.raw_cached = {raw_ws[f"{c}{HEADER_ROW}"].value: raw_ws[f"{c}{r}"].value for c in formula_cols}
    cl = values[CHANGELOG_SHEET]
    changelog = []
    for r in range(HEADER_ROW, cl.max_row + 1):
        row = [cl.cell(r, c).value for c in range(1, 6)]
        if row[3]:
            changelog.append(dict(zip(("date", "samples", "category", "content", "impact"), row, strict=True)))
    return samples, changelog


# ---------------------------------------------------------------------------
# 公式重算与校验
# ---------------------------------------------------------------------------


def _stage(total: int) -> str:
    return "早期信号" if total <= 2 else ("中期确认信号" if total <= 5 else "高风险")


def recompute_summary(s: LegacySample) -> dict[str, Any]:
    """按 Excel 公式的含义重算 J、K、L、R、S、T、U 列。"""
    out: dict[str, Any] = {}
    for ver, (tot, stage, det) in (("v2-M", ("J", "K", "L")), ("v3-R1", ("R", "S", "T"))):
        sc = s.scores[ver]
        if any(v is None for v in sc):
            out.update({tot: None, stage: None, det: None})
            continue
        total = sum(sc)  # type: ignore[arg-type]
        out[tot] = total
        out[stage] = _stage(total)
        price, breadth, vix, _, credit = sc
        out[det] = "是" if price == 2 and breadth == 2 and (credit == 2 or vix == 2) else "否"
    if out["J"] is None or out["R"] is None:
        out["U"] = None
    else:
        out["U"] = f"v2-M:{'是' if out['J'] >= 3 else '否'} / v3-R1:{'是' if out['R'] >= 3 else '否'}"
    return out


def recompute_raw(raw: dict[str, Any]) -> dict[str, float | None]:
    """重算原始数据表的公式列（AH..AX），单位与 Excel 相同（百分比为小数）。"""

    def diff(a: str, b: str) -> float | None:
        return None if raw.get(a) is None or raw.get(b) is None else raw[a] - raw[b]

    def pct(k: str) -> float | None:
        return None if raw.get(k) is None else raw[k] / 100

    out: dict[str, float | None] = {}
    for sym in ("SPY", "QQQ", "RSP"):
        out[f"{sym} 收盘−MA20"] = diff(f"{sym}.close", f"{sym}.ma20")
        out[f"{sym} 收盘−MA50"] = diff(f"{sym}.close", f"{sym}.ma50")
        out[f"{sym} MA5−MA50"] = diff(f"{sym}.ma5", f"{sym}.ma50")
    out["SPY 收盘−MA200"] = diff("SPY.close", "SPY.ma200")
    f, w = pct("F"), pct("W")
    out["L=min(F,W)"] = None if f is None or w is None else min(f, w)
    out["F−F5 (百分点)"] = None if raw.get("F5") is None or f is None else (f - pct("F5")) * 100  # type: ignore[operator]
    out["W−W5 (百分点)"] = None if raw.get("W5") is None or w is None else (w - pct("W5")) * 100  # type: ignore[operator]
    v, v5 = raw.get("VIX"), raw.get("VIX_T5")
    out["VIX g"] = None if v is None or not v5 else v / v5 - 1
    y, h, y20, o1, o6 = pct("y"), pct("H"), pct("y_t20"), pct("oas_o1"), pct("oas_o6")
    out["y−H (bp)"] = None if y is None or h is None else round((y - h) * 10000)
    out["Δy (bp)"] = None if y is None or y20 is None else round((y - y20) * 10000)
    out["ΔOAS (bp)"] = None if o1 is None or o6 is None else round((o1 - o6) * 10000)
    return out


def formula_checks(s: LegacySample) -> list[str]:
    problems = []
    for col, want in recompute_summary(s).items():
        got = s.cached.get(col)
        got = None if got == "" else got
        if want != got:
            problems.append(f"{s.base_date} 样本汇总 {col} 列：程序重算 {want!r}，Excel 缓存 {got!r}")
    for name, want in recompute_raw(s.raw).items():
        # Excel 表头的空格写法与程序可能不同，按去空格后的名称匹配
        got_key = next((k for k in s.raw_cached if k and k.replace(" ", "") == name.replace(" ", "")), None)
        got = s.raw_cached.get(got_key) if got_key else None
        got = None if got == "" else got
        if (want is None) != (got is None) or (want is not None and abs(float(want) - float(got)) > 1e-6):
            problems.append(f"{s.base_date} 原始数据 {name}：程序重算 {want!r}，Excel 缓存 {got!r}")
    return problems


# ---------------------------------------------------------------------------
# 导入
# ---------------------------------------------------------------------------

LEGACY_CALC = "截图时代的录入分数（ChatGPT 打分、经 Claude 核查），见 legacy.json"


def legacy_results(s: LegacySample) -> list[Any]:
    results = []
    for ver in VERSIONS:
        dims = [DimensionScore(n, v, (v,), (), LEGACY_CALC) for n, v in zip(DIM_NAMES, s.scores[ver], strict=True)]
        results.append(assemble(ver, *dims, notes=[s.texts["remark"]] if s.texts.get("remark") else []))
    return results


def legacy_metrics(raw: dict[str, Any]) -> dict[str, Any]:
    m = {k: v for k, v in raw.items() if k != "data_note"}
    f, w = raw.get("F"), raw.get("W")
    m["L"] = None if f is None or w is None else min(f, w)
    for ver in ("v3r1", "v2m"):   # 截图时代两版 O1/O6 相同（Excel 日期对照表说明）
        m[f"oas_o1_{ver}"], m[f"oas_o6_{ver}"] = raw.get("oas_o1"), raw.get("oas_o6")
        o1, o6 = raw.get("oas_o1"), raw.get("oas_o6")
        m[f"doas_{ver}_bp"] = None if o1 is None or o6 is None else round((o1 - o6) * 100, 6)
    return m


def _existing_legacy_run(paths: StoragePaths, base: dt.date, sha: str) -> str | None:
    for run_id in list_runs(paths, MARKET, RISK_SCORING, base):
        meta = read_json(paths.run_dir(MARKET, RISK_SCORING, base, run_id) / RUN_FILES["meta"])
        if meta.get("data_source_type") == "screenshot" and meta.get("legacy", {}).get("sha256") == sha:
            return run_id
    return None


def _latest_program_run(paths: StoragePaths, base: dt.date) -> tuple[str, Path] | None:
    found = None
    for run_id in list_runs(paths, MARKET, RISK_SCORING, base):
        d = paths.run_dir(MARKET, RISK_SCORING, base, run_id)
        meta = read_json(d / RUN_FILES["meta"])
        if meta.get("data_source_type") == "api" and meta.get("status") != "failed":
            found = (run_id, d)
    return found


def compare_with_program(
    s: LegacySample, legacy_key: str, program: tuple[str, Path] | None
) -> tuple[list[Review], list[str], list[str]]:
    """截图记录与程序记录对照：返回 (复核记录, 分数差异, 数值差异)。"""
    if program is None:
        review = Review(MARKET, s.base_date, "program", "截图与程序对照", "没有程序运行记录可供对照", "", legacy_key)
        return [review], [], []
    run_id, d = program
    prog_key = f"{MARKET}/{RISK_SCORING}/{s.base_date}/{run_id}"
    scores = read_json(d / RUN_FILES["scores"])
    score_diffs = []
    for r in scores["results"]:
        prog = [r[k]["score"] for k in ("price", "breadth", "vix", "rates", "credit")]
        legacy = s.scores[r["version"]]
        for name, a, b in zip(DIM_NAMES, legacy, prog, strict=True):
            if a != b:
                score_diffs.append(f"{s.base_date} {r['version']} {name}：截图 {a}，程序 {b}")
    from market_risk.metrics import metrics_from_snapshot_json

    pm = metrics_from_snapshot_json(read_json(d / RUN_FILES["snapshot"]))
    metric_diffs = []
    for key, v in s.raw.items():
        if key.endswith("_date") or key in ("data_note", "H_date") or v is None:
            continue
        pkey = {"oas_o1": "oas_o1_v3r1", "oas_o6": "oas_o6_v3r1"}.get(key, key)
        p = pm.get(pkey)
        if p is None:
            continue
        tol = PRICE_TOL if "." in key else 0.005
        if abs(float(v) - float(p)) > tol + 1e-9:
            metric_diffs.append(f"{s.base_date} {key}：截图 {v}，程序 {round(float(p), 4)}")
    content = "五项分数（两个版本）一致" if not score_diffs else "；".join(score_diffs)
    if metric_diffs:
        content += "；数值差异：" + "；".join(metric_diffs)
    else:
        content += "；原始读数在容差内一致（价格与均线 ≤0.02，其余 ≤0.005）"
    impact = "分数一致" if not score_diffs else "【需用户判断】维度分数不同，未自动处理"
    return [Review(MARKET, s.base_date, "program", "截图与程序对照", content, impact, legacy_key, prog_key)], \
        score_diffs, metric_diffs


def import_legacy(
    excel: Path,
    paths: StoragePaths,
    git: GitInfo,
    now: dt.datetime | None = None,
) -> LegacyImportResult:
    samples, changelog = read_excel(excel)
    sha = hashlib.sha256(excel.read_bytes()).hexdigest()
    result = LegacyImportResult()
    reviews: list[Review] = []
    for s in samples:
        if not s.complete:
            result.skipped.append(f"{s.base_date}（{s.status or '未填写'}）")
            continue
        result.formula_checks += formula_checks(s)
        run_id = _existing_legacy_run(paths, s.base_date, sha)
        if run_id is None:
            run_id, run_dir, created = new_run(paths, MARKET, RISK_SCORING, s.base_date, git, now)
            meta = base_meta(run_id, created, MARKET, RISK_SCORING, s.base_date, "backtest", git,
                             data_source_type="screenshot")
            meta.update(status=STATUS_COMPLETE, legacy={
                "excel": excel.resolve().relative_to(paths.root.resolve()).as_posix()
                if paths.root.resolve() in excel.resolve().parents else excel.as_posix(),
                "sha256": sha, "row": s.row, "seq": s.seq,
            })
            write_json(run_dir / RUN_FILES["meta"], meta)
            write_json(run_dir / RUN_FILES["scores"], {"results": legacy_results(s), "near_threshold": []})
            write_json(run_dir / RUN_FILES["metrics"], legacy_metrics(s.raw))
            write_json(run_dir / RUN_FILES["legacy"], {"summary": {**s.texts, "status": s.status,
                                                                  "scores": s.scores, "cached": s.cached},
                                                      "raw": s.raw, "raw_cached": s.raw_cached})
            (run_dir / RUN_FILES["summary"]).write_text(_legacy_summary(s), encoding="utf-8")
            result.imported.append((s.base_date, run_id))
        else:
            result.skipped.append(f"{s.base_date}（已导入：{run_id}）")
        set_official(paths, MARKET, RISK_SCORING, s.base_date, run_id, "import-legacy", reviewed=True, now=now)
        legacy_key = f"{MARKET}/{RISK_SCORING}/{s.base_date}/{run_id}"
        rv, sd, md = compare_with_program(s, legacy_key, _latest_program_run(paths, s.base_date))
        reviews += rv
        result.score_differences += sd
        result.metric_differences += md

    by_seq = {s.seq: s.base_date for s in samples}
    for row in changelog:
        refs = [int(x) for x in re.findall(r"样本(\d+)", str(row["samples"] or ""))]
        dates = [by_seq[n] for n in refs if n in by_seq] or [None]
        if "及以后" in str(row["samples"] or ""):
            dates = [None]
        for d in dates:
            reviews.append(Review(MARKET, d, "legacy_changelog", f"变更记录：{row['category']}",
                                  f"[{row['date']}]（{row['samples']}）{row['content']}", str(row["impact"] or "")))
            result.changelog_rows += 1
    add_reviews(paths, reviews, now,
                replace=lambda r: r["reviewer"] in ("legacy_changelog",) or r["category"] == "截图与程序对照")
    return result


def _legacy_summary(s: LegacySample) -> str:
    lines = [f"# {s.base_date} 截图时代记录（import-legacy）", "",
             "| 维度 | v2-M | v3-R1 |", "|---|---|---|"]
    for i, n in enumerate(DIM_NAMES):
        lines.append(f"| {n} | {s.scores['v2-M'][i]} | {s.scores['v3-R1'][i]} |")
    lines += ["", f"贴近门槛：{s.texts.get('near_threshold') or '无'}", "",
              f"备注：{s.texts.get('remark') or '无'}", "", f"ChatGPT 会话：{s.texts.get('session') or '无'}"]
    return "\n".join(lines) + "\n"
