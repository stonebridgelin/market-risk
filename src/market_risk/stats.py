"""回测统计（SOP 9.4 / docs/STORAGE.md 第6节）。

只统计正式记录；分别列出已复核与未复核样本数；样本量不足时结论为"不确定"，不得修改标准来凑结论。
标签只在此处读取（评分与 prompt 不读取，STORAGE 第5节）。
"""

from __future__ import annotations

import datetime as dt
import json
import math
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from market_risk import calendar as mcal
from market_risk.outcomes import Outcome, effective_outcomes, read_outcomes
from market_risk.storage import db
from market_risk.storage.paths import StoragePaths

VERSIONS = ("v2-M", "v3-R1")
DIMENSIONS = ("价格", "广度", "VIX", "利率", "信用")
MIN_EVENTS = 40
MIN_DETERMINABLE = 0.90
Z95 = 1.959963984540054


@dataclass
class Sample:
    base_date: dt.date
    run_key: str
    data_source_type: str
    reviewed: bool
    totals: dict[str, dict]                  # version → totals 行
    dims: dict[str, dict[str, dict]]         # version → 维度 → dimension_scores 行
    outcome: Outcome | None


@dataclass
class VersionStats:
    version: str
    n: int = 0
    determinable: int = 0
    alerts: int = 0
    labeled: int = 0
    hits: int = 0
    false_alarms: int = 0
    misses: int = 0
    silences: int = 0
    unknown_alert: list[dt.date] = field(default_factory=list)
    lead_days: list[int] = field(default_factory=list)

    @property
    def events(self) -> int:
        return self.hits + self.misses

    @property
    def non_events(self) -> int:
        return self.false_alarms + self.silences


def _rate(a: int, b: int) -> float | None:
    return None if b == 0 else a / b


def wilson(k: int, n: int) -> tuple[float, float] | None:
    if n == 0:
        return None
    p = k / n
    denom = 1 + Z95**2 / n
    centre = (p + Z95**2 / (2 * n)) / denom
    half = Z95 * math.sqrt(p * (1 - p) / n + Z95**2 / (4 * n * n)) / denom
    return centre - half, centre + half


def newcombe_diff(k1: int, n1: int, k2: int, n2: int) -> tuple[float, float] | None:
    """两个比例之差 p1−p2 的 95% 区间（Newcombe 方法10，按独立样本近似；两版本为同一样本，偏保守）。"""
    a, b = wilson(k1, n1), wilson(k2, n2)
    if a is None or b is None:
        return None
    p1, p2 = k1 / n1, k2 / n2
    d = p1 - p2
    lo = d - math.sqrt((p1 - a[0]) ** 2 + (b[1] - p2) ** 2)
    hi = d + math.sqrt((a[1] - p1) ** 2 + (p2 - b[0]) ** 2)
    return lo, hi


def load_samples(paths: StoragePaths, framework: str, start: dt.date | None, end: dt.date | None,
                 url: str) -> list[Sample]:
    """正式记录及其分数（经 storage.db 读取）与结果标签。"""
    outcomes = effective_outcomes(read_outcomes(paths.outcomes_csv))
    samples = []
    for r in db.official_runs(url, framework):
        d = r["base_date"]
        if (start and d < start) or (end and d > end):
            continue
        samples.append(Sample(d, r["run_key"], r["data_source_type"], bool(r["reviewed"]),
                              db.totals_by_version(url, r["run_key"]), db.dimensions_by_version(url, r["run_key"]),
                              outcomes.get((r["subject"], d))))
    return samples


def version_stats(samples: list[Sample], version: str) -> VersionStats:
    st = VersionStats(version)
    for s in samples:
        t = s.totals.get(version)
        if t is None:
            continue
        st.n += 1
        alert = t["alert"]
        if alert == "未知":
            st.unknown_alert.append(s.base_date)
            continue
        st.determinable += 1
        st.alerts += alert == "是"
        if s.outcome is None:
            continue
        st.labeled += 1
        event = s.outcome.is_event
        if alert == "是" and event:
            st.hits += 1
            if s.outcome.event_date:
                st.lead_days.append(len(mcal.stock_trading_days(s.base_date, s.outcome.event_date)) - 1)
        elif alert == "是":
            st.false_alarms += 1
        elif event:
            st.misses += 1
        else:
            st.silences += 1
    return st


def modification_for(dimension: str, v2: dict, v3: dict) -> str:
    """两版本在某维度分数不同时，对应 SOP 7.4 的哪一处修改。"""
    if dimension == "价格":
        trig = json.loads(v2.get("triggered_conditions") or "[]")
        if "2分(b)" in trig:
            return "修改2：三环节移出计分"
        return "修改1：MA50 持续性过滤" if trig else "修改1或修改2（截图记录无触发条件明细）"
    return {"广度": "修改3：广度第二个2分条件", "信用": "修改4：OAS 时效与计数"}.get(
        dimension, "（VIX、利率两版本规则相同，不应出现差异，请检查）")


def _pct(v: float | None) -> str:
    return "-" if v is None else f"{v:.2f}%"


def _near_text(samples: list[Sample]) -> str:
    near = [s.base_date for s in samples if s.outcome is not None and s.outcome.near_event]
    return f"{len(near)} 个" + (f"（{'、'.join(map(str, near))}）" if near else "")


def _fmt(p: float | None) -> str:
    return "-" if p is None else f"{p * 100:.1f}%"


def render_markdown(samples: list[Sample], stats: dict[str, VersionStats], generated_at: str,
                    start: dt.date | None, end: dt.date | None) -> str:
    reviewed = sum(s.reviewed for s in samples)
    labeled = sum(s.outcome is not None for s in samples)
    lines = [
        "# 回测统计（SOP 9.4）", "",
        f"生成时间（UTC）：{generated_at}；范围：{start or '最早'} 至 {end or '最新'}；只统计正式记录。", "",
        f"- 正式记录样本数：{len(samples)}（已复核 {reviewed}，未复核 {len(samples) - reviewed}）",
        f"- 有结果标签的样本：{labeled}；无标签（结果窗口未结束或未计算）：{len(samples) - labeled}",
        f"- 数据来源：截图 {sum(s.data_source_type == 'screenshot' for s in samples)}，"
        f"程序 {sum(s.data_source_type == 'api' for s in samples)}",
        f"- 接近事件（仅参考，不改变风险事件定义）：{_near_text(samples)}", "",
        "## 指标", "",
        "| 指标 | " + " | ".join(VERSIONS) + " |", "|---|---|---|",
    ]

    def row(name: str, f: object) -> None:
        lines.append(f"| {name} | " + " | ".join(str(f(stats[v])) for v in VERSIONS) + " |")  # type: ignore[operator]

    row("可确定预警状态的样本 / 全部", lambda s: f"{s.determinable}/{s.n}")
    row("数据完整率", lambda s: _fmt(_rate(s.determinable, s.n)))
    row("报警比例（预警/可确定）", lambda s: _fmt(_rate(s.alerts, s.determinable)))
    row("命中 / 误报 / 漏报 / 正确静默", lambda s: f"{s.hits} / {s.false_alarms} / {s.misses} / {s.silences}")
    row("命中率（命中/事件）", lambda s: _fmt(_rate(s.hits, s.events)))
    row("误报率（误报/非事件）", lambda s: _fmt(_rate(s.false_alarms, s.non_events)))
    row("预警准确率（命中/预警）", lambda s: _fmt(_rate(s.hits, s.hits + s.false_alarms)))
    row("正确静默率（静默/非事件）", lambda s: _fmt(_rate(s.silences, s.non_events)))
    row("命中样本提前量（交易日，均值）",
        lambda s: "-" if not s.lead_days else f"{sum(s.lead_days) / len(s.lead_days):.1f}")
    row("状态未知的样本", lambda s: "、".join(map(str, s.unknown_alert)) or "无")

    v2, v3 = stats["v2-M"], stats["v3-R1"]
    fa = newcombe_diff(v3.false_alarms, v3.non_events, v2.false_alarms, v2.non_events)
    hit = newcombe_diff(v3.hits, v3.events, v2.hits, v2.events)
    checks = [
        ("至少90%样本可确定预警状态",
         all(s.n and s.determinable / s.n >= MIN_DETERMINABLE for s in (v2, v3))),
        (f"至少{MIN_EVENTS}个事件和{MIN_EVENTS}个非事件样本",
         min(v2.events, v2.non_events, v3.events, v3.non_events) >= MIN_EVENTS),
        ("v3-R1 相对 v2-M 误报率差的95%区间上限<0", fa is not None and fa[1] < 0),
        ("命中率差的95%区间下限>−5个百分点", hit is not None and hit[0] > -0.05),
    ]
    enough = checks[0][1] and checks[1][1]
    lines += ["", "## 通过标准（均需满足）", "", "| 标准 | 结果 |", "|---|---|"]
    lines += [f"| {name} | {'满足' if ok else '不满足'} |" for name, ok in checks]
    fa_txt = "-" if fa is None else f"{fa[0] * 100:+.1f} 至 {fa[1] * 100:+.1f} 个百分点"
    hit_txt = "-" if hit is None else f"{hit[0] * 100:+.1f} 至 {hit[1] * 100:+.1f} 个百分点"
    lines += ["", f"- 误报率差（v3-R1 − v2-M）95%区间：{fa_txt}", f"- 命中率差（v3-R1 − v2-M）95%区间：{hit_txt}",
              "- 区间按 Newcombe 方法（独立样本近似）计算，仅供参考。", ""]
    if not enough:
        conclusion = "**不确定**：样本量不足（SOP 9.4），不给结论，不得修改标准来凑结论。"
    elif all(ok for _, ok in checks):
        conclusion = "满足全部通过标准。"
    else:
        conclusion = "样本量已足够，但未满足全部通过标准。"
    lines += ["## 结论", "", conclusion, "", "## 两个版本的差异", ""]
    diff_lines = []
    for s in samples:
        for dim in DIMENSIONS:
            a, b = s.dims.get("v2-M", {}).get(dim), s.dims.get("v3-R1", {}).get(dim)
            if a and b and a["score"] != b["score"]:
                diff_lines.append(f"| {s.base_date} | {dim} | {a['score']} | {b['score']} | "
                                  f"{modification_for(dim, a, b)} |")
    if diff_lines:
        lines += ["| 基准日 | 维度 | v2-M | v3-R1 | 对应 SOP 7.4 |", "|---|---|---|---|---|", *diff_lines]
    else:
        lines.append("所有样本两个版本的各维度分数相同。")
    lines += ["", "## 样本明细", "",
              "最低收盘跌幅：以基准日收盘价为起点（风险事件口径）；最大收盘跌幅：窗口内最高收盘价到其后最低收盘价（辅助）。",
              "",
              "| 基准日 | 来源 | 已复核 | v2-M 总分 | v2-M 预警 | v3-R1 总分 | v3-R1 预警 | 风险事件 | "
              "标普500 最低/最大收盘跌幅 | QQQ 最低/最大收盘跌幅 | 接近事件 |",
              "|---|---|---|---|---|---|---|---|---|---|---|"]
    for s in samples:
        t2, t3 = s.totals.get("v2-M", {}), s.totals.get("v3-R1", {})

        def tot(t: dict) -> str:
            return str(t.get("total")) if t.get("total") is not None else f"{t.get('total_min')}–{t.get('total_max')}"

        o = s.outcome
        ev = "-" if o is None else ("是" if o.is_event else "否")
        spx = "-" if o is None else f"{o.spx_min_close_drawdown:.2f}% / {_pct(o.spx_max_drawdown)}"
        qqq = "-" if o is None else f"{o.qqq_min_close_drawdown:.2f}% / {_pct(o.qqq_max_drawdown)}"
        near = "-" if o is None else ("是" if o.near_event else "否")
        lines.append(f"| {s.base_date} | {'截图' if s.data_source_type == 'screenshot' else '程序'} | "
                     f"{'是' if s.reviewed else '否'} | {tot(t2)} | {t2.get('alert')} | {tot(t3)} | "
                     f"{t3.get('alert')} | {ev} | {spx} | {qqq} | {near} |")
    return "\n".join(lines) + "\n"


HISTORY_HEADERS = ["序号", "基准日", "星期", "状态", "v2-M 价格", "v2-M 广度", "v2-M VIX", "v2-M 利率", "v2-M 信用",
                   "v2-M 总分", "v2-M 阶段", "v2-M 明确恶化", "v3-R1 价格", "v3-R1 广度", "v3-R1 VIX", "v3-R1 利率",
                   "v3-R1 信用", "v3-R1 总分", "v3-R1 阶段", "v3-R1 明确恶化", "预警(总分≥3)", "结果窗口",
                   "风险事件(是/否)", "贴近门槛的读数", "Remark：参数变化与数据问题", "数据来源",
                   "标普500最低收盘跌幅(%)", "QQQ最低收盘跌幅(%)", "标普500最大收盘跌幅(%)", "QQQ最大收盘跌幅(%)",
                   "接近事件(仅参考)"]


def write_history_xlsx(path: Path, samples: list[Sample], url: str) -> None:
    """样本汇总表（列与旧 Excel 的"样本汇总"相同，最后一列改为数据来源）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "样本汇总"
    ws["A1"] = "美股大盘风险评分 回测样本汇总（v2-M 与 v3-R1 并行，程序生成）"
    ws["A2"] = "由 market-risk stats 生成，只含正式记录；数值为程序计算或导入的录入值。风险事件为结果标签。"
    for c, h in enumerate(HISTORY_HEADERS, start=1):
        cell = ws.cell(4, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for i, s in enumerate(samples, start=1):
        r = 4 + i
        near = db.near_threshold_for(url, s.run_key)
        values = [i, s.base_date, f"周{'一二三四五六日'[s.base_date.weekday()]}",
                  "已复核" if s.reviewed else "未复核"]
        for v in VERSIONS:
            t = s.totals.get(v, {})
            values += [s.dims.get(v, {}).get(d, {}).get("score") for d in DIMENSIONS]
            values += [t.get("total"), t.get("stage"), t.get("clear_deterioration")]
        values.append(f"v2-M:{s.totals.get('v2-M', {}).get('alert')} / v3-R1:{s.totals.get('v3-R1', {}).get('alert')}")
        w0, w1 = mcal.shift_trading_days(s.base_date, 1), mcal.shift_trading_days(s.base_date, 20)
        values.append(f"{w0}至{w1}")
        values.append("" if s.outcome is None else ("是" if s.outcome.is_event else "否"))
        values.append("；".join(f"{n['item']} {float(n['gap']):+.2f}{n['unit']}" for n in near
                                if n["gap"] is not None))
        notes = json.loads(s.totals.get("v2-M", {}).get("notes") or "[]")
        values.append("；".join(notes))
        values.append("截图" if s.data_source_type == "screenshot" else "程序")
        o = s.outcome
        values += ([o.spx_min_close_drawdown, o.qqq_min_close_drawdown, o.spx_max_drawdown, o.qqq_max_drawdown,
                    "是" if o.near_event else "否"] if o is not None else [None] * 5)
        for c, v in enumerate(values, start=1):
            ws.cell(r, c, v)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def run_stats(paths: StoragePaths, url: str, framework: str = "risk_scoring", start: dt.date | None = None,
              end: dt.date | None = None, now: dt.datetime | None = None) -> tuple[str, list[Sample]]:
    """重建数据库后统计，写 reports/backtest_stats.md 与 reports/backtest_history.xlsx。"""
    db.rebuild(paths, url)
    samples = load_samples(paths, framework, start, end, url)
    stats = {v: version_stats(samples, v) for v in VERSIONS}
    generated = (now or dt.datetime.now(dt.UTC)).astimezone(dt.UTC).isoformat(timespec="seconds")
    text = render_markdown(samples, stats, generated, start, end)
    paths.backtest_stats_md.parent.mkdir(parents=True, exist_ok=True)
    paths.backtest_stats_md.write_text(text, encoding="utf-8")
    write_history_xlsx(paths.backtest_history_xlsx, samples, url)
    return text, samples
