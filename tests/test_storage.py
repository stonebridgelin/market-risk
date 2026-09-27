"""阶段4.5 存储测试（docs/STORAGE.md 第9节验收）。只写入临时目录。"""

from __future__ import annotations

import ast
import dataclasses
import datetime as dt
import json
import shutil
from pathlib import Path

import openpyxl
import pytest
from conftest import load_sample_raw

from market_risk.config import PROJECT_ROOT, load_settings
from market_risk.legacy import import_legacy, read_excel, recompute_summary
from market_risk.materials import MaterialError, Review, add_material, add_reviews, list_materials, read_reviews
from market_risk.models import BreadthReading
from market_risk.outcomes import (
    Outcome,
    OutcomeError,
    compute_outcome,
    effective_outcomes,
    read_outcomes,
    record_outcome,
)
from market_risk.pipeline import run_scoring
from market_risk.stats import run_stats
from market_risk.storage import db, runs
from market_risk.storage.paths import MARKET, RISK_SCORING, StoragePaths

D = dt.date
SETTINGS = load_settings()
GIT = runs.GitInfo("b" * 40, False)
EXCEL = PROJECT_ROOT / "data" / "legacy" / "backtest_record_legacy.xlsx"
SAMPLE_BREADTH = {
    "2025-08-29": {D(2025, 8, 29): (63.22, 67.59)},
    "2025-09-26": {D(2025, 9, 26): (54.67, 51.09)},
    "2025-10-31": {D(2025, 10, 31): (40.15, 38.56), D(2025, 10, 24): (52.88, 57.65)},
    "2025-11-28": {D(2025, 11, 28): (58.44, 76.73)},
}


@pytest.fixture()
def paths(tmp_path):
    return StoragePaths(tmp_path)


def program_runs(paths: StoragePaths, now: dt.datetime | None = None) -> dict[str, str]:
    out = {}
    for sample, readings in SAMPLE_BREADTH.items():
        breadth = {d: BreadthReading(d, f, w) for d, (f, w) in readings.items()}
        out[sample] = run_scoring(load_sample_raw(sample, breadth), SETTINGS, paths, GIT, now=now).run_id
    return out


def copy_excel(paths: StoragePaths) -> Path:
    target = paths.legacy_excel
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(EXCEL, target)
    return target


# ---------------------------------------------------------------------------
# 数据库：删除后重建，内容一致；同一基准日只有一个正式记录
# ---------------------------------------------------------------------------


def test_rebuild_db_is_reproducible(paths):
    program_runs(paths)
    run_scoring(load_sample_raw("2025-11-28", {D(2025, 11, 28): BreadthReading(D(2025, 11, 28), 58.44, 76.73)}),
                SETTINGS, paths, GIT)   # 同一基准日第二次运行
    record_outcome(paths.outcomes_csv, Outcome("MARKET", D(2025, 8, 29), D(2025, 9, 2), D(2025, 9, 29),
                                               -1.2, -2.0, False, None, "manual", "2026-09-27T00:00:00+00:00"))
    note = paths.root / "note.md"
    note.write_text("笔记", encoding="utf-8")
    add_material(paths, "MARKET", D(2025, 11, 28), "notes", note)
    add_reviews(paths, [Review("MARKET", D(2025, 11, 28), "claude", "核查", "无问题")])

    first = db.dump(db.rebuild(paths))
    paths.db_path.unlink()
    shutil.rmtree(paths.db_path.parent)
    second = db.dump(db.rebuild(paths))
    assert first == second
    assert len(first["runs"]) == 5 and len(first["outcomes"]) == 1
    assert len(first["materials"]) == 1 and len(first["reviews"]) == 1
    assert len(first["dimension_scores"]) == 5 * 2 * 5 and len(first["totals"]) == 5 * 2
    officials = db.query(paths.db_path, "SELECT base_date, COUNT(*) n FROM runs WHERE is_official=1 GROUP BY base_date")
    assert {r["base_date"]: r["n"] for r in officials} == {s: 1 for s in SAMPLE_BREADTH}
    metrics = {r["key"]: r["value"] for r in db.query(
        paths.db_path, "SELECT key, value FROM metrics m JOIN runs r USING(run_key) WHERE r.base_date='2025-10-31'")}
    assert metrics["SPY.close"] == 682.06 and metrics["L"] == 38.56 and metrics["doas_v3r1_bp"] == pytest.approx(-11)


def test_alert_status():
    assert db.alert_status(3, 3, 3) == "是" and db.alert_status(2, 2, 2) == "否"
    assert db.alert_status(None, 3, 5) == "是" and db.alert_status(None, 0, 2) == "否"
    assert db.alert_status(None, 2, 4) == "未知"


# ---------------------------------------------------------------------------
# 结果标签与隔离（STORAGE 第5节）
# ---------------------------------------------------------------------------

ISOLATED = [
    *sorted((PROJECT_ROOT / "src" / "market_risk" / "scoring").glob("*.py")),
    PROJECT_ROOT / "src" / "market_risk" / "report.py",
    PROJECT_ROOT / "src" / "market_risk" / "data" / "snapshot.py",
    PROJECT_ROOT / "templates" / "prompt_backtest.md.j2",
]


@pytest.mark.parametrize("path", ISOLATED, ids=lambda p: p.name)
def test_label_isolation(path):
    """评分、prompt 生成与快照不得引用结果标签（文件、表、模块）。"""
    text = path.read_text(encoding="utf-8")
    for forbidden in ("outcomes", "outcomes_csv", "market_risk.outcomes", "is_event", "drawdown",
                      "market_risk.stats", "FROM outcomes"):
        assert forbidden not in text, f"{path.name} 引用了 {forbidden}"
    if path.suffix == ".py":
        tree = ast.parse(text)
        modules = [n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) and n.module]
        modules += [a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names]
        assert not any(m.endswith(("outcomes", "stats", "storage.db", "legacy")) for m in modules), modules


def _closes(base: dt.date, path: list[float]) -> dict[dt.date, float]:
    from market_risk import calendar as mcal

    days = [base, *mcal.stock_trading_days(base + dt.timedelta(days=1), base + dt.timedelta(days=40))[:20]]
    return dict(zip(days, path, strict=True))


def test_compute_outcome():
    base = D(2025, 10, 31)
    spx = _closes(base, [100.0] + [99.0] * 5 + [94.9] + [98.0] * 14)     # 第6个交易日跌 5.1%
    qqq = _closes(base, [100.0] + [97.0] * 20)
    o = compute_outcome(base, spx, qqq, today_new_york=D(2026, 1, 1))
    assert o.is_event and o.spx_min_close_drawdown == pytest.approx(-5.1)
    assert o.event_date == sorted(spx)[6] and (o.window_start, o.window_end) == (D(2025, 11, 3), D(2025, 12, 1))
    calm = compute_outcome(base, _closes(base, [100.0] + [95.01] * 20), _closes(base, [100.0] + [93.01] * 20),
                           D(2026, 1, 1))
    assert not calm.is_event and calm.event_date is None
    with pytest.raises(OutcomeError, match="才结束"):
        compute_outcome(base, spx, qqq, today_new_york=D(2025, 12, 1))   # 窗口最后一天当天不得计算
    del qqq[D(2025, 11, 14)]
    with pytest.raises(OutcomeError, match="缺少收盘价"):
        compute_outcome(base, spx, qqq, D(2026, 1, 1))


def test_manual_and_computed_outcomes(paths):
    manual = Outcome("MARKET", D(2025, 10, 31), D(2025, 11, 3), D(2025, 12, 1), -4.0, -6.0, False, None,
                     "manual", "t")
    computed = dataclasses.replace(manual, spx_min_close_drawdown=-5.2, is_event=True, source="computed")
    assert record_outcome(paths.outcomes_csv, manual) == []
    diffs = record_outcome(paths.outcomes_csv, computed)
    assert len(diffs) == 1 and "标签不一致" in diffs[0]
    rows = read_outcomes(paths.outcomes_csv)
    assert len(rows) == 2
    assert effective_outcomes(rows)[("MARKET", D(2025, 10, 31))].source == "manual"
    record_outcome(paths.outcomes_csv, computed)      # 同一来源重复写入不重复
    assert len(read_outcomes(paths.outcomes_csv)) == 2


# ---------------------------------------------------------------------------
# 资料
# ---------------------------------------------------------------------------


def test_materials(paths, tmp_path):
    f = tmp_path / "chatgpt.md"
    f.write_text("第一版", encoding="utf-8")
    t1 = add_material(paths, "market", D(2025, 11, 28), "chatgpt_response", f, "ChatGPT", "样本4 打分")
    assert t1 == paths.materials_dir("MARKET", D(2025, 11, 28)) / "chatgpt.md"
    f.write_text("第二版", encoding="utf-8")
    t2 = add_material(paths, "MARKET", D(2025, 11, 28), "chatgpt_response", f)
    assert t2.name == "chatgpt_2.md" and t1.read_text("utf-8") == "第一版"   # 不覆盖
    shot = tmp_path / "spy.png"
    shot.write_bytes(b"png")
    assert "screenshots" in add_material(paths, "MARKET", D(2025, 11, 28), "screenshot", shot).parts
    assert len(list_materials(paths, "MARKET", D(2025, 11, 1))) == 3
    assert list_materials(paths, "MARKET", D(2025, 12, 1)) == []
    with pytest.raises(MaterialError):
        add_material(paths, "MARKET", D(2025, 11, 28), "video", f)
    with pytest.raises(MaterialError):
        add_material(paths, "MARKET", D(2025, 11, 28), "notes", tmp_path / "missing.md")


# ---------------------------------------------------------------------------
# import-legacy
# ---------------------------------------------------------------------------


def test_read_excel_and_recompute_formulas():
    samples, changelog = read_excel(EXCEL)
    done = [s for s in samples if s.complete]
    assert [s.base_date for s in done] == [D(2025, 8, 29), D(2025, 9, 26), D(2025, 10, 31), D(2025, 11, 28)]
    assert done[2].scores == {"v2-M": [1, 2, 0, 0, 0], "v3-R1": [1, 2, 0, 0, 0]}
    assert recompute_summary(done[2])["U"] == "v2-M:是 / v3-R1:是"
    assert done[2].raw["F"] == pytest.approx(40.15) and done[2].raw["SPY.ma50"] == 662.17
    assert len(changelog) == 12


def test_import_legacy_end_to_end(paths):
    prog = program_runs(paths)
    excel = copy_excel(paths)
    now = dt.datetime(2026, 9, 27, 3, 0, 0, tzinfo=dt.UTC)
    result = import_legacy(excel, paths, GIT, now=now)
    assert [d for d, _ in result.imported] == [D(2025, 8, 29), D(2025, 9, 26), D(2025, 10, 31), D(2025, 11, 28)]
    assert result.formula_checks == []          # 公式重算与 Excel 缓存值一致
    assert result.score_differences == []       # 截图与程序分数一致
    assert result.metric_differences == []
    assert len(result.skipped) == 9             # 待做的样本5至13
    for d, run_id in result.imported:
        pointer = runs.read_official(paths, MARKET, RISK_SCORING, d)
        assert pointer == {**pointer, "run_id": run_id, "set_by": "import-legacy", "reviewed": True}
        meta = json.loads((paths.run_dir(MARKET, RISK_SCORING, d, run_id) / "meta.json").read_text("utf-8"))
        assert meta["data_source_type"] == "screenshot" and meta["legacy"]["sha256"]
        # 程序记录作为对照保留
        assert prog[d.isoformat()] in runs.list_runs(paths, MARKET, RISK_SCORING, d)
    reviews = read_reviews(paths)
    compare = [r for r in reviews if r["category"] == "截图与程序对照"]
    assert len(compare) == 4 and all(r["impact"] == "分数一致" for r in compare)
    assert all(r["other_run_key"].endswith(prog[r["base_date"]]) for r in compare)
    assert sum(r["reviewer"] == "legacy_changelog" for r in reviews) == result.changelog_rows >= 12

    # 重复导入：不新建运行目录，不重复复核记录
    again = import_legacy(excel, paths, GIT, now=now)
    assert again.imported == [] and any("已导入" in s for s in again.skipped)
    assert len(read_reviews(paths)) == len(reviews)


def test_import_legacy_reports_score_difference(paths):
    prog = program_runs(paths)
    scores_path = paths.run_dir(MARKET, RISK_SCORING, D(2025, 10, 31), prog["2025-10-31"]) / "scores.json"
    scores = json.loads(scores_path.read_text("utf-8"))
    scores["results"][1]["price"]["score"] = 2          # 人为制造差异
    scores_path.write_text(json.dumps(scores, ensure_ascii=False), encoding="utf-8")
    result = import_legacy(copy_excel(paths), paths, GIT)
    assert result.score_differences == ["2025-10-31 v3-R1 价格：截图 1，程序 2"]
    row = next(r for r in read_reviews(paths) if r["base_date"] == "2025-10-31" and r["category"] == "截图与程序对照")
    assert "需用户判断" in row["impact"]


def test_import_legacy_without_program_runs(paths):
    result = import_legacy(copy_excel(paths), paths, GIT)
    assert len(result.imported) == 4 and result.score_differences == []
    assert any("没有程序运行记录" in r["content"] for r in read_reviews(paths))


# ---------------------------------------------------------------------------
# stats
# ---------------------------------------------------------------------------


def test_stats_with_legacy_samples(paths):
    program_runs(paths)
    import_legacy(copy_excel(paths), paths, GIT)
    record_outcome(paths.outcomes_csv, Outcome("MARKET", D(2025, 10, 31), D(2025, 11, 3), D(2025, 12, 1),
                                               -5.3, -7.9, True, D(2025, 11, 20), "manual", "t"))
    record_outcome(paths.outcomes_csv, Outcome("MARKET", D(2025, 8, 29), D(2025, 9, 2), D(2025, 9, 29),
                                               -1.0, -1.5, False, None, "manual", "t"))
    text, samples = run_stats(paths, now=dt.datetime(2026, 9, 27, tzinfo=dt.UTC))
    assert len(samples) == 4 and all(s.data_source_type == "screenshot" for s in samples)
    assert "正式记录样本数：4（已复核 4，未复核 0）" in text
    assert "| 命中 / 误报 / 漏报 / 正确静默 | 1 / 0 / 0 / 1 | 1 / 0 / 0 / 1 |" in text
    assert "**不确定**" in text
    assert "所有样本两个版本的各维度分数相同" in text
    assert "| 命中样本提前量（交易日，均值） | 14.0 | 14.0 |" in text
    assert paths.backtest_stats_md.read_text("utf-8") == text

    # 样本汇总表的分数与旧 Excel 一致
    old = openpyxl.load_workbook(EXCEL, data_only=True)["样本汇总"]
    new = openpyxl.load_workbook(paths.backtest_history_xlsx)["样本汇总"]
    for i in range(4):
        old_row = [old.cell(5 + i, c).value for c in range(5, 21)]
        new_row = [new.cell(5 + i, c).value for c in range(5, 21)]
        assert new_row == old_row
    assert new.cell(7, 23).value == "是" and new.cell(5, 26).value == "截图"


def test_stats_unreviewed_and_mixed(paths):
    program_runs(paths)     # 自动设为正式记录，reviewed=false
    text, _ = run_stats(paths)
    assert "已复核 0，未复核 4" in text and "程序 4" in text
